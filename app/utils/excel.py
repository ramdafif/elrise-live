from datetime import date, datetime
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from defusedxml.ElementTree import iterparse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.utils.validation import ValidationError

IMPORT_HEADERS = ["employee_number", "name", "email", "department", "job_title", "phone", "join_date", "office_id", "password"]


def workbook_bytes(headers, rows, title):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="102F40")
    for values in rows:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cells in sheet.columns:
        sheet.column_dimensions[cells[0].column_letter].width = min(42, max(18, len(str(cells[0].value)) + 4))
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def read_employees(upload):
    if not upload or not upload.filename.lower().endswith(".xlsx"):
        raise ValidationError("Gunakan file Excel .xlsx dari template Elrise.")
    content = upload.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise ValidationError("File Excel maksimal 5 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 200 or sum(entry.file_size for entry in entries) > 20 * 1024 * 1024:
                raise ValidationError("File Excel terlalu besar setelah diekstrak.")
            if any(entry.flag_bits & 1 or "externalLinks/" in entry.filename or "vbaProject" in entry.filename for entry in entries):
                raise ValidationError("File terenkripsi, macro, atau tautan eksternal tidak diizinkan.")
            for entry in entries:
                if entry.filename.startswith("xl/worksheets/") and entry.filename.endswith(".xml"):
                    with archive.open(entry) as stream:
                        for event, element in iterparse(stream, events=("end",)):
                            if element.tag.endswith("}row"):
                                if int(element.attrib.get("r", "0")) > 501 and any(child.tag.endswith("}c") and len(child) for child in element):
                                    raise ValidationError("Import maksimal 500 baris; data di luar rentang tidak diabaikan.")
                                element.clear()
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) != 1:
                raise ValidationError("Template harus memiliki satu sheet.")
            sheet = workbook.active
            sheet.reset_dimensions()
            lines = sheet.iter_rows(max_row=502, max_col=len(IMPORT_HEADERS) + 1)
            header = [cell.value for cell in next(lines)]
            if header[:len(IMPORT_HEADERS)] != IMPORT_HEADERS or header[-1] is not None:
                raise ValidationError("Nama/urutan kolom tidak sesuai template.")
            rows = []
            for line_number, cells in enumerate(lines, start=2):
                values = [cell.value for cell in cells]
                if all(value is None for value in values):
                    continue
                if line_number > 501 or values[-1] is not None:
                    raise ValidationError("Import maksimal 500 baris, tanpa kolom tambahan.")
                if any(cell.data_type == "f" for cell in cells):
                    raise ValidationError(f"Baris {line_number}: formula tidak diizinkan.")
                data = dict(zip(IMPORT_HEADERS, values))
                for key, value in data.items():
                    if isinstance(value, datetime):
                        data[key] = value.date().isoformat()
                    elif isinstance(value, date):
                        data[key] = value.isoformat()
                    else:
                        data[key] = "" if value is None else str(value)
                rows.append((line_number, data))
            if not rows:
                raise ValidationError("File belum berisi data karyawan.")
            return rows
        finally:
            workbook.close()
    except ValidationError:
        raise
    except Exception as error:
        raise ValidationError("File Excel rusak atau tidak dapat dibaca.") from error
