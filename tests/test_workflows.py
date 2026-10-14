from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.db import execute, query
from app.utils.excel import IMPORT_HEADERS, workbook_bytes
from app.utils.validation import now
from tests.conftest import TEST_PASSWORD, attendance_payload, consent, form_post, login, submit


def employee_data(office_id, suffix="NEW"):
    return {"name": "New Team Member " + suffix, "employee_number": "EMP-" + suffix,
        "email": suffix.lower() + "@team.test", "department": "Engineering", "job_title": "Engineer",
        "phone": "", "join_date": now().date().isoformat(), "office_id": str(office_id), "password": "Initial-Password!2026"}


def test_master_flow_superadmin_to_employee_to_report(app, selfie):
    owner = app.test_client()
    assert login(owner, "owner@elrise.test").status_code == 302
    today = now().date()
    response = form_post(owner, "/superadmin/clients/new", {"name": "PT End To End", "code": "E2E", "email": "hr@e2e.test", "phone": "", "address": "Jakarta", "employee_limit": "10", "service_start": str(today), "service_end": str(today + timedelta(days=365))})
    assert response.status_code == 302
    with app.app_context():
        company = query("SELECT * FROM clients WHERE code='E2E'", one=True)
        assert not company["is_active"]
    client_id = company["id"]
    initial = "Bootstrap-Password!2026"
    assert form_post(owner, f"/superadmin/clients/{client_id}/admins", {"name": "HR E2E", "email": "hr@e2e.test", "password": initial}).status_code == 302
    admin = app.test_client()
    assert login(admin, "hr@e2e.test", initial).status_code == 401
    assert form_post(owner, f"/superadmin/clients/{client_id}/toggle").status_code == 302
    assert login(admin, "hr@e2e.test", initial).location.endswith("/password")
    admin_password = "New-Admin-Password!2026"
    assert form_post(admin, "/password", {"current_password": initial, "new_password": admin_password, "confirm_password": admin_password}).status_code == 302
    assert login(admin, "hr@e2e.test", admin_password).location.endswith("/admin/")
    assert form_post(admin, "/admin/company", {"name": "PT End To End", "email": "hr@e2e.test", "address": "Kantor Jakarta", "phone": "021123456", "photo_retention_days": "14"}).status_code == 302
    office_data = {"name": "Kantor E2E", "address": "Jakarta", "latitude": "-6.2", "longitude": "106.816666", "radius_m": "100", "max_accuracy_m": "50", "start_time": "09:00", "end_time": "17:00", "is_active": "1", "working_days": ["1", "2", "3", "4", "5", "6", "7"]}
    assert form_post(admin, "/admin/offices/new", office_data).status_code == 302
    with app.app_context():
        office_id = query("SELECT id FROM office_locations WHERE client_id=%s", (client_id,), one=True)["id"]
    data = employee_data(office_id, "E2E")
    assert form_post(admin, "/admin/employees/new", data).status_code == 302
    employee = app.test_client()
    assert login(employee, data["email"], data["password"]).location.endswith("/password")
    employee_password = "New-Employee-Password!2026"
    assert form_post(employee, "/password", {"current_password": data["password"], "new_password": employee_password, "confirm_password": employee_password}).status_code == 302
    assert login(employee, data["email"], employee_password).location.endswith("/employee/")
    consent(employee)
    assert employee.get("/employee/attendance").status_code == 200
    check_in = submit(employee, attendance_payload(employee, selfie))
    assert check_in.status_code == 200
    attendance_id = check_in.get_json()["attendance_id"]
    assert submit(employee, attendance_payload(employee, selfie)).status_code == 200
    assert employee.get("/employee/history").status_code == 200
    detail = admin.get(f"/admin/attendance/{attendance_id}")
    assert detail.status_code == 200 and "Selfie check-in" in detail.get_data(as_text=True)
    assert admin.get(f"/employee/photos/{attendance_id}/in").status_code == 200
    exported = admin.get("/admin/attendance/export.xlsx?start=" + str(today) + "&end=" + str(today))
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.data))
    rows = list(workbook.active.values)
    assert rows[1][2] == data["name"] and rows[1][6] and rows[1][7]
    assert "Bima Adiputra" not in str(rows) and "Rani Prameswari" not in str(rows)


def test_employee_import_real_accounts_and_excel_formula_safety(app, records):
    admin = app.test_client()
    login(admin, "admin@nusantara.test")
    template = admin.get("/admin/employees/template.xlsx")
    assert template.status_code == 200
    assert list(load_workbook(BytesIO(template.data)).active.values)[0] == tuple(IMPORT_HEADERS)
    data = employee_data(records["employee"]["office_id"], "IMPORTED")
    data["name"] = "=HYPERLINK(\"http://example.test\",\"unsafe\")"
    content = workbook_bytes(IMPORT_HEADERS, [[data[key] for key in IMPORT_HEADERS]], "Karyawan")
    response = form_post(admin, "/admin/employees/import", {"file": (content, "employees.xlsx")}, content_type="multipart/form-data")
    assert response.status_code == 302, response.get_data(as_text=True)
    assert login(app.test_client(), data["email"], data["password"]).location.endswith("/password")
    exported = admin.get("/admin/attendance/export.xlsx")
    assert exported.status_code == 200
    sheet = load_workbook(BytesIO(exported.data)).active
    names = [row[2] for row in list(sheet.rows)[1:]]
    imported_name = next(cell for cell in names if cell.value == data["name"])
    assert imported_name.data_type == "s"
    assert not any(cell.data_type == "f" for row in sheet for cell in row)


def test_import_is_atomic_and_rejects_formula_and_oversized_row_count(app, records):
    admin = app.test_client()
    login(admin, "admin@nusantara.test")
    first = employee_data(records["employee"]["office_id"], "FIRST")
    duplicate = employee_data(records["employee"]["office_id"], "DUPLICATE")
    duplicate["email"] = "rani@nusantara.test"
    content = workbook_bytes(IMPORT_HEADERS, [[data[key] for key in IMPORT_HEADERS] for data in [first, duplicate]], "Karyawan")
    assert form_post(admin, "/admin/employees/import", {"file": (content, "employees.xlsx")}).status_code == 400
    with app.app_context():
        assert not query("SELECT id FROM users WHERE email=%s", (first["email"],), one=True)
        assert query("SELECT COUNT(*) AS total FROM employees", one=True)["total"] == 2
    book = load_workbook(workbook_bytes(IMPORT_HEADERS, [[first[key] for key in IMPORT_HEADERS]], "Karyawan"))
    book.active["B2"] = "=2+2"
    formula_file = BytesIO()
    book.save(formula_file)
    formula_file.seek(0)
    response = form_post(admin, "/admin/employees/import", {"file": (formula_file, "formula.xlsx")})
    assert response.status_code == 400 and "formula" in response.get_data(as_text=True)
    too_many = workbook_bytes(IMPORT_HEADERS, [[first[key] for key in IMPORT_HEADERS]] * 501, "Karyawan")
    response = form_post(admin, "/admin/employees/import", {"file": (too_many, "too-many.xlsx")})
    assert response.status_code == 400 and "500" in response.get_data(as_text=True)


def test_capacity_limit_and_concurrent_creation(app, records):
    with app.app_context():
        execute("UPDATE clients SET employee_limit=2 WHERE id=%s", (records["client"]["id"],))
    first, second = app.test_client(), app.test_client()
    login(first, "admin@nusantara.test")
    login(second, "admin@nusantara.test")
    requests = [(first, employee_data(records["employee"]["office_id"], "ONE")), (second, employee_data(records["employee"]["office_id"], "TWO"))]
    with ThreadPoolExecutor(max_workers=2) as executor:
        responses = list(executor.map(lambda item: form_post(item[0], "/admin/employees/new", item[1]), requests))
    assert sorted(response.status_code for response in responses) == [302, 400]
    with app.app_context():
        assert query("SELECT COUNT(*) AS total FROM employees WHERE client_id=%s AND is_active=1", (records["client"]["id"],), one=True)["total"] == 2


def test_employee_and_office_crud_and_reset_sessions(app, records):
    admin = app.test_client()
    login(admin, "admin@nusantara.test")
    data = employee_data(records["employee"]["office_id"], "CRUD")
    assert form_post(admin, "/admin/employees/new", data).status_code == 302
    with app.app_context():
        employee = query("SELECT * FROM employees WHERE employee_number='EMP-CRUD'", one=True)
    employee_id = employee["id"]
    data.update(name="Updated Name", is_active="1")
    assert form_post(admin, f"/admin/employees/{employee_id}/edit", data).status_code == 302
    with app.app_context():
        assert query("SELECT name FROM employees WHERE id=%s", (employee_id,), one=True)["name"] == "Updated Name"
    assert form_post(admin, f"/admin/employees/{employee_id}/delete").status_code == 302
    with app.app_context():
        assert not query("SELECT id FROM users WHERE id=%s", (employee["user_id"],), one=True)
    browser = app.test_client()
    login(browser, "rani@nusantara.test")
    assert form_post(admin, f"/admin/employees/{records['employee']['id']}/reset", {"password": "Reset-Password!2026"}).status_code == 302
    assert browser.get("/employee/").status_code == 302
    assert login(browser, "rani@nusantara.test", "Reset-Password!2026").location.endswith("/password")
    assert browser.get("/employee/").location.endswith("/password")
    assert form_post(admin, f"/admin/offices/{records['employee']['office_id']}/delete").status_code == 400


def test_leave_create_approve_and_prevent_attendance(app, selfie):
    employee, admin = app.test_client(), app.test_client()
    login(employee, "rani@nusantara.test")
    login(admin, "admin@nusantara.test")
    today = str(now().date())
    data = {"kind": "cuti", "start_date": today, "end_date": today, "reason": "Personal day"}
    assert form_post(employee, "/employee/leaves", data).status_code == 302
    assert form_post(employee, "/employee/leaves", data).status_code == 400
    with app.app_context():
        leave_id = query("SELECT id FROM leave_requests", one=True)["id"]
    assert admin.get("/admin/leaves").status_code == 200
    assert form_post(admin, f"/admin/leaves/{leave_id}/review", {"status": "approved", "admin_note": "Enjoy your day"}).status_code == 302
    assert form_post(admin, f"/admin/leaves/{leave_id}/review", {"status": "rejected"}).status_code == 409
    consent(employee)
    response = submit(employee, attendance_payload(employee, selfie))
    assert response.status_code == 400 and "disetujui" in response.get_json()["message"]
    assert "Enjoy your day" in employee.get("/employee/leaves").get_data(as_text=True)


def test_leave_cancel_and_cross_tenant_review(app):
    other, admin = app.test_client(), app.test_client()
    login(other, "bima@cakrawala.test")
    login(admin, "admin@nusantara.test")
    tomorrow = str(now().date() + timedelta(days=1))
    data = {"kind": "izin", "start_date": tomorrow, "end_date": tomorrow, "reason": "Other tenant private reason"}
    assert form_post(other, "/employee/leaves", data).status_code == 302
    with app.app_context():
        leave_id = query("SELECT id FROM leave_requests", one=True)["id"]
    assert "Other tenant private reason" not in admin.get("/admin/leaves").get_data(as_text=True)
    assert form_post(admin, f"/admin/leaves/{leave_id}/review", {"status": "approved"}).status_code == 404
    assert form_post(other, f"/employee/leaves/{leave_id}/cancel").status_code == 302
    assert form_post(other, f"/employee/leaves/{leave_id}/cancel").status_code == 404


def test_schema_initialization_refuses_foreign_app_tables(app):
    with app.app_context():
        execute("CREATE TABLE foreign_app (id INT PRIMARY KEY)")
    try:
        result = app.test_cli_runner().invoke(args=["init-db"])
        assert result.exit_code != 0 and "aplikasi lain" in result.output
        with app.app_context():
            assert query("SELECT COUNT(*) AS total FROM users", one=True)["total"] == 5
    finally:
        with app.app_context():
            execute("DROP TABLE foreign_app")


def test_superadmin_edit_client_and_reset_only_matching_admin(app, records):
    owner, admin = app.test_client(), app.test_client()
    login(owner, "owner@elrise.test")
    login(admin, "admin@nusantara.test")
    company = records["client"]
    changed = {"name": "PT Updated", "code": company["code"], "email": company["email"], "phone": "021123", "address": "Updated address", "employee_limit": "100",
        "service_start": str(company["service_start"]), "service_end": str(company["service_end"]), "photo_retention_days": "14", "is_active": "1"}
    assert form_post(owner, f"/superadmin/clients/{company['id']}/edit", changed).status_code == 302
    with app.app_context():
        current = query("SELECT * FROM clients WHERE id=%s", (company["id"],), one=True)
        assert current["name"] == "PT Updated" and current["photo_retention_days"] == 14
        own_admin = query("SELECT id FROM users WHERE email='admin@nusantara.test'", one=True)["id"]
        other_admin = query("SELECT id FROM users WHERE email='admin@cakrawala.test'", one=True)["id"]
    assert form_post(owner, f"/superadmin/clients/{company['id']}/admins/{other_admin}/reset", {"password": "Reset-Admin!2026"}).status_code == 404
    assert form_post(owner, f"/superadmin/clients/{company['id']}/admins/{own_admin}/reset", {"password": "Reset-Admin!2026"}).status_code == 302
    assert admin.get("/admin/").location.endswith("/login")
    assert login(admin, "admin@nusantara.test", "Reset-Admin!2026").location.endswith("/password")


@pytest.mark.parametrize("address,paths", [
    ("owner@elrise.test", ["/superadmin/clients", "/superadmin/clients/new", "/password"]),
    ("admin@nusantara.test", ["/admin/company", "/admin/employees", "/admin/employees/new", "/admin/employees/import", "/admin/offices", "/admin/offices/new", "/admin/attendance", "/admin/leaves"]),
    ("rani@nusantara.test", ["/employee/", "/employee/privacy", "/employee/history", "/employee/leaves", "/employee/profile", "/password"]),
])
def test_all_main_pages_render_real_data(app, address, paths):
    browser = app.test_client()
    login(browser, address)
    for path in paths:
        response = browser.get(path)
        assert response.status_code == 200, path + "\n" + response.get_data(as_text=True)
        assert 'lang="id"' in response.get_data(as_text=True)
